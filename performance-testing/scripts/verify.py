from pathlib import Path
import json,hashlib,fitz
from pptx import Presentation
P=Path(__file__).resolve().parents[1];R=P.parent
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
proof={}
proof['analytics_equal']=json.loads((P/'results/before/analytics-snapshot.json').read_text(encoding='utf-8'))==json.loads((P/'results/after/analytics-snapshot.json').read_text(encoding='utf-8'))
assert proof['analytics_equal']
for stem in ['load-1','load-2','load-3','stress-3','soak-3']:
    b=json.loads((P/'results/before'/(stem+'-meta.json')).read_text());a=json.loads((P/'results/after'/(stem+'-meta.json')).read_text())
    assert b['script_sha256']==a['script_sha256']==digest(P/'scripts/scenario.js')
    assert b['controller_sha256']==digest(P/'baseline/TeacherCourseController.php')
    assert a['controller_sha256']==digest(P/'baseline/TeacherCourseController.after.php')
    assert b['exit_code'] in [0,99] and a['exit_code'] in [0,99]
proof['same_test_script_and_correct_controllers']=True
assert digest(R/'app/Http/Controllers/TeacherCourseController.php')==digest(P/'baseline/TeacherCourseController.after.php')
proof['database_unchanged']=digest(P/'runtime/performance.sqlite')==json.loads((P/'reports/environment.json').read_text())['database_sha256']
assert proof['database_unchanged']
d=fitz.open(P/'documentation/KS_Semestralny_projekt_Tema14_Mykhailo_Adamenko.pdf');proof['pdf_pages']=len(d)
assert 8<=len(d)<=10
for pg in d:
    assert abs(pg.rect.width-595.276)<1 and abs(pg.rect.height-841.89)<1
    assert '\ufffd' not in pg.get_text()
    assert 'TODO' not in pg.get_text()
assert 'Použité zdroje' in d[-1].get_text()
slides=Presentation(P/'documentation/KS_Tema14_Prezentacia_Mykhailo_Adamenko.pptx');proof['presentation_slides']=len(slides.slides)
assert len(slides.slides)==12
for s in slides.slides:
    assert s.notes_slide.notes_text_frame.text.strip()
    for shape in s.shapes:
        assert shape.left>=0 and shape.top>=0
        assert shape.left+shape.width<=slides.slide_width+1000
        assert shape.top+shape.height<=slides.slide_height+1000
proof['frontend_build_exit']=json.loads((P/'results/after/frontend-build-status.json').read_text())['returncode'];assert proof['frontend_build_exit']==0
for phase in ['before','after']:
    c=json.loads((P/'results'/phase/'controls.json').read_text())
    assert c['metrics']['control_errors']['values']['rate']==0
    assert c['metrics']['checks']['values']['passes']==20
proof['control_checks_passed']=40
(P/'reports/verification.json').write_text(json.dumps(proof,indent=2),encoding='utf-8')
print(json.dumps(proof,indent=2))
