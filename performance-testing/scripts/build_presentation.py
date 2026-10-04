"""Editable visual presentation generated exclusively from measured results."""
from pathlib import Path
import json,statistics,math,shutil
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
P=Path(__file__).resolve().parents[1];D=P/'documentation'
raw=json.loads((P/'reports/comparison.json').read_text(encoding='utf-8'));details=raw['details']
def row(ph,key):return next(r for r in raw['rows'] if r['phase']==ph and r['test']==key)
def f(v,n=1):return f'{v:.{n}f}'.replace('.',',')
def readj(p):
 b=p.read_bytes();return json.loads(b.decode('utf-16' if b[:2] in [b'\xff\xfe',b'\xfe\xff'] else 'utf-8-sig'))
sections=readj(P/'results/before/profile-sections.json')['samples'][1:]
sec={k:statistics.median(x['sections'][k] for x in sections) for k in sections[0]['sections']}
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
NAVY='101D30';PANEL='1B2B41';WHITE='F4F7FA';INK='172B42';MUTED='627387';TEAL='19B7A4';ORANGE='EB9960';RED='CA5C69';LINE='DBE3EB';PALE='E8F5F2'
def rgb(c):return RGBColor.from_string(c)
def rect(s,x,y,w,h,fill,outline=None,round=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if outline:sh.line.color.rgb=rgb(outline)
 else:sh.line.fill.background()
 if round:sh.adjustments[0]=.12
 return sh
def txt(s,x,y,w,h,text,size=18,color=INK,bold=False,font='Aptos',align=None):
 if color==TEAL and s.background.fill.fore_color.rgb==rgb(WHITE):color='078476'
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(5)
  if align is not None:p.alignment=align
 return sh
def line(s,x1,y1,x2,y2,color=LINE,width=1):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width);return sh
def circle(s,x,y,d,fill):
 sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.fill.background();return sh
def pill(s,x,y,w,label,color=TEAL,dark=False):
 rect(s,x,y,w,.34,PANEL if dark else PALE,round=True);txt(s,x+.12,y+.055,w-.24,.24,label,11,color,True)
def note(s,text,seconds):s.notes_slide.notes_text_frame.text=f'Odporúčaný čas: {seconds} s.\n\n{text}'
notes=[]
def slide(k,title,sub='',dark=False):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(NAVY if dark else WHITE)
 fg=WHITE if dark else INK
 rect(s,.55,.48,.12,.34,TEAL);txt(s,.82,.47,11,.3,k.upper(),11,TEAL,True)
 txt(s,.65,.96,12.05,.73,title,33,fg,True)
 if sub:txt(s,.67,1.72,12,.48,sub,16,'AEBED0' if dark else MUTED)
 line(s,.65,7.04,12.67,7.04,PANEL if dark else LINE)
 txt(s,.67,7.14,10,.2,'MYKHAILO ADAMENKO   /   FEI STU   /   KVALITA SOFTVÉROVÝCH SYSTÉMOV',9,'91A4B9' if dark else MUTED)
 txt(s,12.1,7.09,.55,.3,f'{len(prs.slides):02d} / 12',10,TEAL,True)
 notes.append((title,sub,s));return s
def metric(s,x,y,w,value,label,detail='',dark=False,color=TEAL):
 rect(s,x,y,w,1.45,PANEL if dark else 'FFFFFF',round=True)
 txt(s,x+.2,y+.14,w-.4,.64,value,34,color,True)
 txt(s,x+.2,y+.86,w-.4,.3,label,15,WHITE if dark else INK,True)
 if detail:txt(s,x+.2,y+1.19,w-.4,.2,detail,10,'AEBED0' if dark else MUTED)
def bar(s,x,y,w,value,maxv,label,fill,fg=INK):
 txt(s,x,y-.01,2.0,.3,label,15,fg)
 rect(s,x+2.0,y+.03,w-2.95,.22,LINE if fg==INK else PANEL,round=True)
 rect(s,x+2.0,y+.03,max(.02,(w-2.95)*value/maxv),.22,fill,round=True)
 txt(s,x+w-.82,y-.02,.85,.32,f(value),14,fg,True)
def plot(s,x,y,w,h,series,labels,maxv,ticks,dark=False):
 fg='AEBED0' if dark else MUTED;grid=PANEL if dark else LINE
 for t in ticks:
  yy=y+h-h*t/maxv;line(s,x,yy,x+w,yy,grid);txt(s,x-.53,yy-.13,.45,.25,f(t,0),10,fg)
 for j,label in enumerate(labels):txt(s,x+j*w/(len(labels)-1)-.38,y+h+.15,.9,.3,str(label),12,fg)
 for vals,col in series:
  for j in range(len(vals)-1):line(s,x+j*w/(len(vals)-1),y+h-h*vals[j]/maxv,x+(j+1)*w/(len(vals)-1),y+h-h*vals[j+1]/maxv,col,2.5)
  for j,v in enumerate(vals):circle(s,x+j*w/(len(vals)-1)-.047,y+h-h*v/maxv-.047,.094,col)
def legend(s,x,y,dark=False):
 rect(s,x,y+.05,.16,.16,ORANGE);txt(s,x+.25,y,1.0,.28,'BEFORE',11,'AEBED0' if dark else MUTED,True)
 rect(s,x+1.4,y+.05,.16,.16,TEAL);txt(s,x+1.65,y,1,.28,'AFTER',11,'AEBED0' if dark else MUTED,True)

# 1 — the topic comes first; the application is introduced later.
s=slide('Semestrálny projekt · Téma 14','Výkonnostné testovanie',dark=True)
txt(s,.75,2.34,7.3,1.5,'Funguje aplikácia dobre,\naj keď ju používa viac ľudí?',34,WHITE,True)
txt(s,.78,4.37,6.8,.85,'Rýchlosť · spoľahlivosť · stabilita',23,'AEBED0')
txt(s,.78,6.25,7.6,.35,'Mykhailo Adamenko  ·  FEI STU  ·  2026/2027',16,'AEBED0')
rect(s,9.0,2.5,3.0,2.15,PANEL,round=True)
rect(s,9.24,2.75,2.52,1.42,NAVY,round=True)
line(s,10.5,4.65,10.5,5.1,TEAL,5);line(s,9.8,5.1,11.2,5.1,TEAL,5)
for j,h in enumerate([.35,.65,.95,1.12]):rect(s,9.57+j*.5,4.01-h,.27,h,TEAL,round=True)
note(s,'Predstavujem tému výkonnostného testovania. Aplikácia môže správne vykonať požadovanú operáciu, no používateľ na ňu môže príliš dlho čakať. Najprv vysvetlím, čo pri výkone sledujeme a ako sa testuje. Potom postup ukážem na aplikácii z mojej bakalárskej práce, ktorá tu slúži ako praktický príklad.',35)

# 2 — motivation, before tools or implementation.
s=slide('01 / Prečo testovať výkon?','Správna odpoveď môže prísť príliš neskoro.')
for x,title,col,many in [(.8,'Jeden používateľ',TEAL,False),(7.0,'Veľa používateľov naraz',ORANGE,True)]:
 rect(s,x,2.2,5.52,3.42,'FFFFFF',LINE,round=True)
 txt(s,x+.3,2.5,4.92,.55,title,24,INK,True)
 for j in range(6 if many else 1):
  xx=x+.4+j*.48;circle(s,xx,3.49,.21,col);rect(s,xx-.035,3.74,.28,.42,col,round=True)
 txt(s,x+3.4,3.64,.48,.5,'→',28,MUTED)
 rect(s,x+4.12,3.4,.9,.75,INK,round=True)
 txt(s,x+.3,4.75,4.9,.45,'Rýchla odpoveď?' if not many else 'Čakanie alebo chyby?',23,INK,True)
txt(s,.83,6.16,11.7,.48,'To, čo funguje pri jednom človeku, treba overiť aj pri väčšej záťaži.',23,TEAL,True)
note(s,'Ide o ilustráciu otázky, nie o výsledok merania. Funkčné testovanie overuje, či aplikácia vracia správne výsledky. Výkonnostné testovanie dopĺňa otázku, ako sa správa pri rôznom počte požiadaviek alebo pri rastúcom objeme údajov. Pomalá odpoveď môže používateľa prinútiť opakovať akciu a zhoršiť jeho skúsenosť.',45)

# 3 — accessible vocabulary.
s=slide('02 / Čo sledujeme?','Tri otázky, ktorým rozumie aj používateľ.')
for j,(num,title,body) in enumerate([('01','Rýchlosť','Ako dlho čakám\nna odpoveď?'),('02','Spoľahlivosť','Dokončí sa moja\npožiadavka bez chyby?'),('03','Stabilita','Funguje aplikácia dobre\naj po dlhšom čase?')]):
 x=.78+j*4.1;rect(s,x,2.4,3.78,3.5,'FFFFFF',LINE,round=True)
 txt(s,x+.28,2.7,2.9,.62,num,36,TEAL,True)
 txt(s,x+.28,3.69,3.25,.5,title,25,INK,True)
 txt(s,x+.28,4.5,3.24,.93,body,21,MUTED)
note(s,'Rýchlosť opisuje čas odpovede. Spoľahlivosť sledujeme cez chyby a kontroly správnosti odpovedí; samotná absencia chybového kódu nestačí. Stabilita znamená, či sa správanie počas testu výrazne nezhoršuje. V dokumentácii je tiež priepustnosť, teda počet vybavených požiadaviek za sekundu. Pri meraní možno sledovať aj využitie procesora a pamäte.',50)

# 4 — types of tests as simple load diagrams.
s=slide('03 / Ako sa výkon testuje?','Každý typ testu odpovedá na inú otázku.')
configs=[('Bežná záťaž','Zvládne očakávanú prevádzku?',[(0,0),(.15,.55),(.85,.55),(1,0)]),('Rastúca záťaž','Kedy začne mať problémy?',[(0,.15),(.25,.15),(.25,.4),(.5,.4),(.5,.65),(.75,.65),(.75,1),(1,1)]),('Dlhšia prevádzka','Nezhoršuje sa časom?',[(0,.55),(1,.55)])]
for j,(title,caption,pts) in enumerate(configs):
 x=.78+j*4.1;rect(s,x,2.3,3.78,3.95,'FFFFFF',LINE,round=True)
 txt(s,x+.25,2.61,3.3,.5,title,23,INK,True)
 line(s,x+.4,4.85,x+3.36,4.85,LINE);line(s,x+.4,3.5,x+.4,4.85,LINE)
 for k in range(len(pts)-1):line(s,x+.4+pts[k][0]*2.96,4.85-pts[k][1]*1.22,x+.4+pts[k+1][0]*2.96,4.85-pts[k+1][1]*1.22,TEAL,3)
 txt(s,x+.28,5.25,3.22,.72,caption,20,MUTED)
txt(s,.83,6.5,11.6,.27,'Schémy znázorňujú počet súčasných používateľov v čase.',13,MUTED)
note(s,'Odborné názvy sú záťažový test (load), stresový test (stress) a vytrvalostný test (endurance alebo soak). Prvý overuje očakávanú prevádzku, druhý hľadá hranice a tretí sleduje vývoj v čase. Diagramy sú len schémy. Neskoršia praktická ukážka používa pri treťom type skrátený trojminútový beh, ktorý nenahrádza dlhodobý test.',55)

# 5 — method before the example.
s=slide('04 / Postup testovania','Najprv cieľ. Potom meranie a porovnanie.')
for j,(num,title) in enumerate([('1','Určiť\nočakávania'),('2','Zmerať\nsprávanie'),('3','Odstrániť\nproblém'),('4','Zopakovať\nrovnaký test')]):
 x=.8+j*3.12;rect(s,x,2.65,2.63,2.66,'FFFFFF',LINE,round=True)
 circle(s,x+.25,2.94,.65,TEAL);txt(s,x+.46,3.015,.3,.4,num,24,WHITE,True)
 txt(s,x+.25,3.95,2.15,1.0,title,22,INK,True)
 if j<3:txt(s,x+2.71,3.62,.38,.5,'→',26,TEAL)
rect(s,.8,5.97,11.96,.69,INK,round=True)
txt(s,1.05,6.12,11.4,.4,'Zrýchlenie má zmysel iba vtedy, keď zostanú výsledky správne.',22,WHITE)
note(s,'Požiadavky stanovíme pred meraním, aby sme ich dodatočne neprispôsobovali výsledkom. Pri opakovaní zachováme rovnaké údaje, scenár a podmienky. Hľadanie príčiny má vychádzať z merania jednotlivých častí spracovania. Po úprave znovu overíme výkon aj správnosť. Tento postup teraz prenesiem do konkrétnej aplikácie.',45)

# 6 — bachelor application is only now introduced as an example.
s=slide('05 / Praktický príklad','Ukážka na aplikácii z bakalárskej práce','Webová aplikácia na online vzdelávanie')
txt(s,.8,2.64,4.0,.5,'Študent',25,TEAL,True)
txt(s,.8,3.2,4.03,.83,'Prechádza lekcie\na rieši testy.',23,INK)
txt(s,.8,4.48,4.0,.5,'Učiteľ',25,TEAL,True)
txt(s,.8,5.04,4.03,.92,'Sleduje priebeh štúdia\na výsledky kurzu.',23,INK)
rect(s,5.16,2.36,7.51,4.4,'FFFFFF',LINE,round=True)
shot=s.shapes.add_picture(str(P/'screenshots/application-analytics.png'),Inches(5.26),Inches(2.48),width=Inches(7.3),height=Inches(4.03))
shot.crop_bottom=1-(4.03/7.3)*(1440/1000)
txt(s,5.35,6.53,7.05,.25,'Obrazovka aplikácie s testovacími údajmi',11,MUTED)
note(s,'Aplikácia vznikla v mojej bakalárskej práci a v tomto projekte slúži ako testovaný príklad. Nejde o prezentáciu celej bakalárskej práce. Vybral som učiteľský prehľad kurzu, pretože spracúva údaje o študentoch, pokroku a výsledkoch. Screenshot je zo skutočnej aplikácie so syntetickými údajmi. Technicky ide o Laravel, Vue/Inertia a SQLite. Meria sa odpoveď servera, nie dokončenie vykreslenia v prehliadači.',55)

# 7 — compact experimental setup and goal.
s=slide('06 / Čo sme overovali?','Zvládne prehľad kurzu rast počtu študentov?')
for j,n in enumerate([50,250,1000]):
 x=.8+j*4.1;metric(s,x,2.38,3.76,str(n).replace('1000','1 000'),'študentov v kurze')
txt(s,.85,4.35,11.65,.5,'Rovnaký test pred úpravou aj po nej.',26,INK,True)
rect(s,.8,5.32,11.96,1.31,INK,round=True)
txt(s,1.07,5.52,11.35,.42,'Cieľ pri 20 súčasných používateľoch: odpoveď do 0,8 s.',23,WHITE,True)
txt(s,1.08,6.08,11.3,.32,'Požadovaná hranica pre 95 zo 100 odpovedí.',16,'AEBED0')
note(s,'Veľkosť kurzu a súčasná záťaž sú dve rozdielne veci. Kurzy obsahovali 50, 250 a 1000 študentov, každý s piatimi pokusmi o test a kurzom s 20 lekciami. Bežná záťaž mala 20 virtuálnych používateľov; 1000 študentov teda neznamená 1000 ľudí online naraz. Automatizované požiadavky vytváral nástroj k6. Cieľ PR-01 bol p95 pod 800 ms pri 20 používateľoch a PR-02 pod 1500 ms pri 50 používateľoch, vždy v stabilnej fáze pre najväčší kurz. Ďalej sa overovali chyby pod 1 %, stabilita a správne výsledky. Testy bežali lokálne, na vývojovom serveri s jedným PHP workerom.',65)

# 8 — plain-language diagnosis, no implementation jargon.
s=slide('07 / Zistený problém','Pri väčšom kurze používateľ dlho čakal.',dark=True)
for j,(label,key) in enumerate([('50 študentov','load-1'),('250 študentov','load-2'),('1 000 študentov','load-3')]):
 v=row('before',key)['p(95)']/1000;y=2.5+j*.9
 txt(s,.84,y,2.9,.45,label,22,WHITE)
 rect(s,3.7,y+.09,6.7*v/21,.28,ORANGE,round=True)
 txt(s,3.88+6.7*v/21,y-.025,1.65,.5,f(v,2)+' s',22,WHITE,True)
txt(s,.87,5.4,11.3,.42,'Hranica času pre 95 zo 100 dokončených odpovedí.',16,'AEBED0')
rect(s,.81,6.09,11.84,.62,PANEL,round=True)
txt(s,1.03,6.22,11.3,.35,'Pri najväčšom kurze sa 11 rozbehnutých opakovaní testu nedokončilo.',18,ORANGE)
note(s,'Porovnávame celé behy pri rovnakej bežnej záťaži a troch veľkostiach kurzu. Ide o p95, nie priemernú odpoveď; päť percent zaznamenaných odpovedí môže byť pomalších. Pri najväčšom kurze bolo 11 prerušených iterácií, preto nulová chybovosť dokončených volaní sama osebe nepotvrdila spoľahlivosť. Lokálny server vybavoval požiadavky jedným workerom, čo vytváralo čakanie vo fronte. Výsledky nevyjadrujú kapacitu produkčného nasadenia.',60)

# 9 — a visual explanation of the change.
s=slide('08 / Čo sme zmenili?','Rovnaké údaje sme spracúvali zbytočne opakovane.')
for x,title,col,body in [(.8,'Pred úpravou',ORANGE,'Tie isté odpovede\nspracované opakovane.'),(7.0,'Po úprave',TEAL,'Odpovede pripravené raz\na znovu použité.')]:
 rect(s,x,2.2,5.52,3.63,'FFFFFF',LINE,round=True)
 txt(s,x+.28,2.49,4.98,.48,title,24,INK,True)
 for j in range(3):
  rect(s,x+.33+j*1.65,3.43,1.42,.62,col if x<1 or j==0 else PALE,round=True)
  txt(s,x+.47+j*1.65,3.56,1.19,.3,'výpočet' if x<1 else ('príprava' if j==0 else 'použitie'),14,INK,True)
 txt(s,x+.29,4.58,4.91,.92,body,23,INK)
txt(s,.86,6.25,11.6,.47,'Výsledky prehľadu zostali zhodné vo všetkých troch kurzoch.',22,TEAL,True)
note(s,'Meranie častí kódu ukázalo problém hlavne v spracovaní údajov v PHP. Opakovane sa dekódovali odpovede pri prechode všetkými otázkami a pokusmi. Po zmene sa odpovede pripravia raz na pokus. Vybrali sme iba potrebné stĺpce, odstránili nepotrebné načítanie používateľa pri výsledku aj ďalšie triedenie už zoradených pokusov. História odpovedí sa zachovala. Porovnanie úplných výstupov analytiky sa zhodovalo pre všetky tri kurzy. Prešlo 51 automatizovaných testov s 244 kontrolami. Samostatný profil kontroléra klesol z mediánu 1028,5 ms na 412,7 ms; ide o iné meranie než čas odpovede pri záťaži.',60)

# 10 — one headline result, with the meaning of the metric preserved.
s=slide('09 / Výsledok','Čakanie sa výrazne skrátilo.',dark=True)
txt(s,.86,2.65,4.1,.45,'Pred úpravou',23,'AEBED0')
txt(s,.82,3.31,4.15,1.1,f(row('before','load-3')['p(95)']/1000)+' s',65,ORANGE,True)
txt(s,5.16,3.39,1.05,.9,'→',51,WHITE)
txt(s,7.0,2.65,4.9,.45,'Po úprave',23,'AEBED0')
txt(s,6.93,3.31,4.6,1.1,f(row('after','load-3')['p(95)']/1000)+' s',65,TEAL,True)
pill(s,7.01,4.68,3.7,'O 56 % KRATŠÍ ČAS',dark=True)
txt(s,.87,5.73,11.5,.75,'Kurz s 1 000 študentmi · 20 súčasných používateľov\nHranica času pre 95 zo 100 dokončených odpovedí.',18,'AEBED0')
note(s,'Hodnoty p95 za celý bežný záťažový beh sú 20,595 s pred úpravou a 9,069 s po úprave, zobrazené zaokrúhlene. Pokles je 55,97 %. Nejde o priemer ani o garanciu času každej odpovede. Dáta a testovací scenár boli rovnaké. Pri vyhodnotení požiadavky používame iba stabilnú fázu: po úprave 9,102 s, čo je stále výrazne nad cieľom 0,8 s. Zlepšenie výkonu preto automaticky neznamená splnenie požiadaviek.',55)

# 11 — measured achievements, scoped to the completed checks.
s=slide('10 / Dosiahnuté výsledky','Rýchlejšie spracovanie. Zachovaná správnosť.')
items=[('Sledovaný čas odozvy','Skrátený o 56 %',TEAL),('Výsledky prehľadu','Zhodné vo všetkých 3 kurzoch',TEAL),('Automatizované testy','51 úspešných testov',TEAL),('Bežná záťaž po úprave','Bez chýb a prerušených behov',TEAL)]
for j,(label,result,col) in enumerate(items):
 y=2.27+j*.84;rect(s,.8,y,11.87,.65,'FFFFFF',LINE,round=True)
 circle(s,1.04,y+.24,.17,col);txt(s,1.42,y+.12,5.35,.41,label,22,INK,True);txt(s,7.08,y+.15,5.16,.38,result,20,col,True)
txt(s,.85,6.18,11.62,.62,'Záťažové porovnanie: lokálny server · 20 súčasných používateľov.',20,MUTED)
note(s,'Na tomto slajde zhrniem dosiahnuté výsledky. Pri kurze s 1000 študentmi a 20 súčasných používateľoch sa sledovaný čas odozvy skrátil približne o 56 %. Ide o hranicu pre 95 zo 100 dokončených odpovedí, vysvetlenú na predchádzajúcom slajde. Celé výstupy prehľadu boli zhodné pred úpravou aj po nej vo všetkých troch veľkostiach kurzu. Prešlo 51 automatizovaných testov s 244 kontrolami. V nameraných behoch bežnej záťaže po úprave neboli zaznamenané chyby ani prerušené iterácie. Tieto výsledky sa vzťahujú na vykonané lokálne merania.',55)

# 12 — topic-level conclusion, then compact sources.
s=slide('11 / Záver','Od merania ku konkrétnemu zlepšeniu.',dark=True)
for j,(title,body) in enumerate([('Zistenie','Našli sme zbytočne\nopakované spracovanie.'),('Úprava','Znížili sme množstvo\nopakovanej práce.'),('Overenie','Potvrdili sme zrýchlenie\na zhodnosť výsledkov.')]):
 x=.78+j*4.1;rect(s,x,2.2,3.78,2.61,PANEL,round=True)
 txt(s,x+.26,2.59,3.23,.49,title,27,TEAL,True)
 txt(s,x+.26,3.46,3.23,.95,body,22,WHITE)
txt(s,.85,5.23,11.6,.66,'Meranie pomohlo vybrať účinnú úpravu a overiť jej výsledok.',23,WHITE)
txt(s,.86,6.24,11.65,.51,'Zdroje: zadanie KS · bakalárska práca autora · dokumentácia Grafana, Laravel a PHP.\nÚplné odkazy sú v poznámkach a písomnej dokumentácii.',12,'AEBED0')
note(s,'Hlavným prínosom projektu je prepojenie merania s konkrétnou úpravou aplikácie. Najprv sme identifikovali opakované spracovanie údajov, potom ho obmedzili a rovnakými testami overili účinok. Výsledkom je namerané zrýchlenie pri zachovaní správnosti prehľadu. Aplikácia z bakalárskej práce tak poslúžila ako praktický príklad využitia výkonnostného testovania. Ďakujem za pozornosť.\n\nPoužité zdroje:\n1. KS_ZS_202627_semestralny_projekt_temy.pdf, téma 14; i-ks_semestralne_temy_hodnotenie.pdf.\n2. Mykhailo Adamenko: bakalárska práca, BP_FINAL.pdf.\n3. Grafana k6: https://grafana.com/docs/k6/latest/using-k6/metrics/ ; https://grafana.com/docs/k6/latest/using-k6/scenarios/ ; https://grafana.com/docs/k6/latest/using-k6/cookies/ .\n4. Laravel: https://laravel.com/docs/12.x/database ; https://laravel.com/docs/12.x/eloquent-mutators .\n5. PHP: https://www.php.net/manual/en/features.commandline.webserver.php .',45)


# Suppress theme shadows on chart lines, symbols and cards.
for sl in prs.slides:
 for shape in sl.shapes:
  for effect in shape._element.xpath('./p:style/a:effectRef'):effect.set('idx','0')
out=D/'KS_Tema14_Prezentacia_Mykhailo_Adamenko.pptx';prs.save(out)
(D/'presentation.md').write_text('\n\n'.join(f'## {i}. {title}\n\n{sub}\n\n{s.notes_slide.notes_text_frame.text}' for i,(title,sub,s) in enumerate(notes,1)),encoding='utf-8')
shutil.copy2(out,P.parent.parent/(out.stem+'_v3'+out.suffix))
print(out)
