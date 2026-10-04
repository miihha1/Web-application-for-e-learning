from pathlib import Path
import json,statistics,html
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.enums import TA_CENTER,TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,Table,TableStyle,Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
P=Path(__file__).resolve().parents[1];D=P/'documentation';D.mkdir(exist_ok=True)
data=json.loads((P/'reports/comparison.json').read_text());rows=data['rows'];details=data['details']
def row(phase,test):return next(x for x in rows if x['phase']==phase and x['test']==test)
def f(v,d=1):return '—' if v is None else f'{v:,.{d}f}'.replace(',',' ').replace('.',',')
def readj(path):
 b=path.read_bytes();return json.loads(b.decode('utf-16' if b[:2] in [b'\xff\xfe',b'\xfe\xff'] else 'utf-8-sig'))
profiles={phase:readj(P/'results'/phase/'profile-large.json')['samples'][1:] for phase in ['before','after']}
prof={phase:{'ms':statistics.median(x['ms'] for x in v),'sql_ms':statistics.median(x['sql_ms'] for x in v),'queries':len(v[-1]['queries'])} for phase,v in profiles.items()}
sections=readj(P/'results/before/profile-sections.json')['samples'][1:]
sec={k:statistics.median(x['sections'][k] for x in sections) for k in sections[0]['sections']}
before=row('before','load-3');after=row('after','load-3');gain=(before['p(95)']-after['p(95)'])/before['p(95)']*100
req=[]
for name,key,limit in [('PR-01','load-3',800),('PR-02','stress-3',1500)]:
 vals=[details[ph+'-'+key]['steady']['p95'] if key=='load-3' else details[ph+'-'+key]['stages']['50']['p95'] for ph in ['before','after']]
 req.append([name]+[f(v)+' ms; '+('áno' if v is not None and v<limit else 'nie') for v in vals])
req.append(['PR-03']+[f(row(ph,'load-3')['error']*100,2)+' %; '+('áno' if row(ph,'load-3')['error']<.01 else 'nie') for ph in ['before','after']])
if row('before','load-3')['interrupted_iterations']:
    req[-1][1]='0 % dokončených; 11 prerušených, nepreukázané'
ratios={}
for ph in ['before','after']:
 t=details[ph+'-soak-3']['thirds'];ratios[ph]=t[-1]['p95']/t[0]['p95'];
req.append(['PR-04']+[f(ratios[ph],3)+'×; '+('áno' if ratios[ph]<=1.25 and all(t['http_error']<.01 for t in details[ph+'-soak-3']['thirds']) else 'nie') for ph in ['before','after']])
req.append(['PR-05','1000 / 5000 / 1000; áno','1000 / 5000 / 1000; áno'])
(P/'reports/requirements.json').write_text(json.dumps(req,ensure_ascii=False,indent=2),encoding='utf-8')
for name,file in [('TR','times.ttf'),('TB','timesbd.ttf'),('TI','timesi.ttf')]:pdfmetrics.registerFont(TTFont(name,'C:/Windows/Fonts/'+file))
pdfmetrics.registerFontFamily('TR',normal='TR',bold='TB',italic='TI',boldItalic='TB')
styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='BodySK',fontName='TR',fontSize=11,leading=15,spaceAfter=9,alignment=TA_JUSTIFY));styles.add(ParagraphStyle(name='TitleSK',fontName='TB',fontSize=17,leading=22,spaceAfter=15));styles.add(ParagraphStyle(name='SubSK',fontName='TB',fontSize=12,leading=16,spaceAfter=7));styles.add(ParagraphStyle(name='CellSK',fontName='TR',fontSize=9,leading=12));styles.add(ParagraphStyle(name='CenterSK',fontName='TR',fontSize=13,leading=19,alignment=TA_CENTER))
story=[];plain=[]
def para(s):story.append(Paragraph(s,styles['BodySK']));plain.append(s)
def heading(s):story.append(Paragraph(s,styles['TitleSK']));plain.append('\n'+s)
def sub(s):story.append(Paragraph(s,styles['SubSK']))
def page():story.append(PageBreak())
def table(head,body,widths=None):
 t=Table([[Paragraph(html.escape(str(c)),styles['CellSK']) for c in r] for r in [head]+body],colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eeeeee')),('LINEBELOW',(0,0),(-1,0),.8,colors.black),('LINEBELOW',(0,-1),(-1,-1),.6,colors.black),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]));story.extend([t,Spacer(1,12)])
def chart(name):story.append(Image(str(P/'reports'/name),width=455,height=193));story.append(Spacer(1,8))
def footer(c,doc):
 if doc.page>1:c.setFont('TR',10);c.drawCentredString(A4[0]/2,28,str(doc.page-1))
story.append(Spacer(1,35))
for s in ['SLOVENSKÁ TECHNICKÁ UNIVERZITA V BRATISLAVE','Fakulta elektrotechniky a informatiky']:story.append(Paragraph(s,styles['CenterSK']))
story.append(Spacer(1,175))
for s in ['<b>Kvalita softvérových systémov</b>','Semestrálny projekt','<b>Téma 14 – Výkonnostné testovanie</b>','E-learning na webovej platforme']:story.append(Paragraph(s,styles['CenterSK']))
story.append(Spacer(1,190));table(['2026','Mykhailo Adamenko'],[],[230,225]);page()
heading('1 Úvod, cieľ a teoretický základ')
para('Cieľom projektu je zistiť, ako sa mení odozva učiteľskej analytiky pri raste kurzu a súbežných požiadavkách, nájsť príčinu oneskorenia a overiť účinok konkrétnej úpravy. Predmetom je existujúca bakalárska aplikácia E-learning na webovej platforme. Téma a forma výstupov vychádzajú zo zadania predmetu.')
para('Výkonnostný test hodnotí čas a spoľahlivosť spracovania pri určenej záťaži. Load test overuje zvolenú bežnú záťaž, stress test ju zvyšuje až po degradáciu a soak test sleduje správanie pri dlhšie trvajúcej stabilnej záťaži. V tomto experimente bol soak skrátený na tri minúty; plná 30-minútová konfigurácia je pripravená, ale nebola meraná.')
para('Čas odozvy sa meria na HTTP úrovni nástrojom k6. Percentil p95 označuje hodnotu, pod ktorou leží približne 95 % zaznamenaných časov. p99 zvýrazňuje pomalý koniec rozdelenia; pri malom počte vzoriek je menej stabilný. Priemer sám osebe nestačí, preto sa ukladajú aj medián, p90 a p99 [1].')
para('Throughput je počet dokončených požiadaviek za sekundu. Chybovosť zahŕňa neúspešné HTTP odpovede a timeouty; osobitná kontrola overuje aj obsah úspešnej odpovede. Uzavretý model VUs obsahuje medzi požiadavkami sekundovú pauzu. Keď server spomalí, VU čaká a ponúkaná záťaž klesá [2].')
sub('Vymedzenie rozsahu')
para('Meraný je GET /teacher/courses/{course}/manage a spracovanie v TeacherCourseController::manage a buildAnalytics. Endpoint vracia zoznam študentov, dokončené lekcie, posledné výsledky, súhrnné počty a najčastejšie chybné otázky. Všetky historické odpovede sú potrebné pre štatistiku otázok; odstránenie histórie by zmenilo význam výstupu.')
para('Projekt používa PHP/Laravel, Vue.js, Inertia.js, TypeScript a Vite. V meraní sa používa SQLite. Zoznam kurzov a odoslanie testu sú doplnkové sériové kontrolné scenáre; hlavné výkonnostné závery sa týkajú iba analytiky. Vykresľovanie stránky v prehliadači nie je súčasťou meranej latencie.')
page();heading('2 Požiadavky na výkonnosť')
para('Požiadavky boli zapísané pred finálnym meraním do requirements/performance_requirements.md. Ide o navrhnuté akceptačné ciele lokálneho experimentu, nie o existujúcu produkčnú SLA. Limity sa po zistení výsledkov nemenili.')
table(['ID','Overiteľné kritérium'],[['PR-01','LARGE, 20 VUs: p95 < 800 ms v stabilnej 30 s fáze load testu.'],['PR-02','LARGE, 50 VUs: p95 < 1500 ms v stabilnej 30 s fáze stress testu.'],['PR-03','Load: HTTP chyby < 1 % a úspešné obsahové kontroly ≥ 99 %.'],['PR-04','Soak 5 VUs / 180 s: p95 poslednej tretiny ≤ 1,25× p95 prvej; chyby každej tretiny < 1 %.'],['PR-05','Pri 1000 študentoch HTTP 200 a počty: 1000 zapísaných, 5000 pokusov, 1000 posledných výsledkov.']],[55,400])
sub('Prečo práve analytika')
para('Náročnosť rastie s počtom študentov, pokusov a záznamov o priebehu štúdia. Pôvodný kontrolér načíta výsledky do Eloquent kolekcie, zoskupí ich podľa študenta a pri každej otázke prechádza všetky pokusy. To je dôvod na meranie, nie automatický dôkaz chyby.')
sub('Zachovanie pôvodného stavu')
para('Experiment prebieha v samostatnej Git vetve performance-testing, založenej na commit-e 80aaabe. Zdrojový archív BEFORE a kópia pôvodného kontroléra umožňujú návrat. Pôvodná databáza používateľa sa nepoužila. Rozdiel implementácie je odovzdaný ako optimization.patch.')
para('Očakávaná funkčnosť zahŕňa zachovanie posledného pokusu podľa času, počtov a priemeru, poradia názvov dokončených lekcií a histórie nesprávnych odpovedí. Kontroly prístupových práv ostávajú bez zmeny. Optimalizácia nesmie zameniť rýchlejšiu neúplnú odpoveď za správny výsledok.')
page();heading('3 Metodika, nástroje a údaje')
para('Použitý bol k6 2.3.0, PHP 8.5.0, Laravel 12.44.0 a inertia-laravel 2.0.18. Počítač: Intel Core i5-11400H, približne 16 GB RAM, Windows. Server aj generátor záťaže bežali lokálne; konfigurácia a hashe sú v reports/environment.json a metadátach behov.')
table(['Dataset','Študenti','Lekcie','Pokusy','Dokončené lekcie'],[['SMALL',50,20,250,667],['MEDIUM',250,20,1250,3333],['LARGE',1000,20,5000,13333]],[85,80,70,80,140])
para('Každý kurz má jeden test s 10 otázkami a 40 možnosťami. Seeder používa seed 202627 a fixné časové značky; v databáze sú všetky tri kurzy. SMALL/MEDIUM sú kontrolou vplyvu objemu, hlavné load/stress/soak scenáre používajú LARGE.')
table(['Test','Profil'],[['Load','5 s nábeh, 30 s pri 20 VUs, 5 s pokles; všetky tri datasety.'],['Stress','10 → 20 → 50 → 100 VUs, 30 s každá fáza + 10 s na dokončenie.'],['Soak','5 VUs, 180 s; porovnanie troch 60 s okien.']],[70,385])
para('Setup vykoná skutočné prihlásenie učiteľa s CSRF tokenom a zahrievací GET. K6 prenesie cookies do VUs a posiela Inertia verziu. Používa sa spoločná session učiteľa, preto test reprezentuje súbežné čítanie jedného dashboardu. Autentifikácia a CSRF neboli vypnuté. Login sa nezapočítava do analytics_ms [3].')
para('PHP development server má vo Windows jeden worker [4]. Jeho fronta je súčasťou pozorovanej odozvy. Timeout HTTP je 30 s. Soak sa začína po reštarte servera s prázdnou frontou. Neúspešný setup po stress teste a prerušený čiastočný soak sú zachované mimo platných výsledkov.')
para('Každá konfigurácia má jeden platný beh BEFORE a AFTER. Ide o prieskumné lokálne porovnanie bez intervalov spoľahlivosti. Presné skripty a datasety zostali rovnaké; časový odstup medzi behmi a bežné procesy systému môžu ovplyvniť absolútne hodnoty.')
page();heading('4 Výsledky BEFORE')
table(['Scenár','p50 ms','p95 ms','p99 ms','RPS','Chyby %'],[[r['test'],f(r['med']),f(r['p(95)']),f(r['p(99)']),f(r['rps'],2),f(r['error']*100,2)] for r in rows if r['phase']=='before'],[80,80,80,80,60,75])
para('Tabuľka uvádza celé behy. RPS je počet dokončených analytických volaní delený celkovým časom behu k6 vrátane setup a grace period. Pri overení PR-01/02 sa používajú len príslušné stabilné časové okná. Úplné avg, medián a percentily sú v CSV/JSON exporte.')
para('Rozdiel medzi SMALL a LARGE ukazuje význam objemu údajov. V odoslanej odpovedi sa prenáša celý zoznam študentov; meranie preto zahŕňa aj serializáciu a prijatie odpovede. Pri vysokej súbežnosti sa k výpočtu pripája čakanie vo fronte jediného workeru.')
sub('Degradácia a neúspešné požiadavky')
para('Load LARGE BEFORE mal 11 prerušených iterácií; nulová chybovosť dokončených volaní preto nedokazuje úplnú spoľahlivosť. Stress BEFORE dosiahol timeouty a 29 prerušených iterácií. Percentily vychádzajú len z HTTP volaní so zaznamenaným ukončením. Fronta pokračovala po ukončení stress klienta, preto následný soak setup zlyhal na prihlásení. Po reštarte sa vykonal nový platný soak.')
para('Samotný počet VUs nie je kapacita aplikácie: 100 čakajúcich VUs neznamená 100 súbežne vykonávaných PHP požiadaviek. Výsledok ukazuje správanie konkrétnej lokálnej konfigurácie. Záver o produkčnom limite by vyžadoval produkčný web server, oddelený generátor a reprezentatívnejšiu prevádzku.')
page();heading('5 Profilovanie a príčina oneskorenia')
para('Po BEFORE meraniach bol kontrolér profilovaný oddelene od HTTP fronty. DB::listen zaznamenal počet a trvanie SQL dotazov [5]. Časovače hrtime rozdelili pôvodnú analytiku na načítanie výsledkov, načítanie progresu, zostavenie študentov a štatistiku otázok. Použila sa instrumentovaná kópia; meraná aplikačná implementácia BEFORE sa nemenila.')
table(['LARGE – CLI profil','BEFORE'],[['Medián času kontroléra + serializácie',f(prof['before']['ms'])+' ms'],['Medián súčtu SQL času',f(prof['before']['sql_ms'])+' ms'],['Počet SQL dotazov',prof['before']['queries']]],[300,155])
table(['Časť pôvodnej analytiky','Medián ms'],[[k,f(v)] for k,v in sec.items()],[300,155])
para('SQL čas je podstatne menší než čas celej požiadavky. Profil preto nepodporuje pridanie indexov ako prvé opatrenie. Zároveň nebol zistený rast počtu dotazov o jeden dotaz na každého študenta v hlavnom endpointe. Zoznam kurzov obsahuje inú potenciálnu N+1 cestu, ale tá nie je predmetom optimalizácie.')
para('Pri výpočte chybných otázok sa opakovane číta atribút answers typu array z Eloquent modelu. Každá otázka znovu spracúva odpovede všetkých pokusov. Pri 10 otázkach a 5000 pokusoch ide o 50 000 prístupov. Výber latestResults navyše triedi už časovo zoradené skupiny. Progres načítava plné modely, hoci sa používajú len identifikátory.')
para('Interpretácia profilu: prioritou je znížiť opakovanú prácu PHP a hydratáciu nepotrebných údajov. História pokusov zostáva zachovaná, pretože ju vyžaduje štatistika otázok. Cache sa nepoužila, aby sa nezaviedla nová povinnosť invalidácie pri zápise výsledkov a progresu.')
page();heading('6 Vykonaná optimalizácia a správnosť')
para('Úprava je sústredená v buildAnalytics. Odstránilo sa nepotrebné eager loading používateľa pri výsledkoch, keďže mená a e-maily sú už dostupné cez enrollments. Výsledky načítavajú len stĺpce používané analytikou. Posledný výsledok sa vyberie ako prvý z už zostupne zoradenej skupiny, bez opakovaného triedenia.')
para('Odpovede pokusov sa dekódujú raz a ďalej sa používajú ako obyčajné polia. Identifikátory vybraných možností sa normalizujú pomocou array_map a sort, čím sa obmedzí tvorba dočasných kolekcií vo vnorenom cykle. Zachováva sa pôvodná práca so skalárnou odpoveďou, prázdnou odpoveďou a duplicitami.')
para('Lesson progress používa query builder a iba user_id/lesson_id. Celé modely, časové značky a ďalšie stĺpce nie sú na zostavenie zoznamu názvov lekcií potrebné. Názvy sa stále vyberajú z lekcií v pôvodnom poradí. Schéma databázy sa nemenila; nevznikla nová cache ani indexová migrácia.')
table(['Overenie','Výsledok'],[['Regresné testy','Posledný pokus, priemer, počty, progres, chybné otázky, prístupové práva.'],['Celý výstup','Analytics JSON na SMALL/MEDIUM/LARGE zhodný BEFORE/AFTER.'],['Aplikačné testy','51 testov, 244 assertions; results/after/functional.txt.'],['Frontend','Úspešný Vite production build; bez zmeny Vue komponentov.']],[125,330])
table(['CLI profil LARGE','BEFORE','AFTER'],[['Medián ms',f(prof['before']['ms']),f(prof['after']['ms'])],['SQL ms',f(prof['before']['sql_ms']),f(prof['after']['sql_ms'])],['SQL dotazy',prof['before']['queries'],prof['after']['queries']]],[215,120,120])
para('Výsledný patch je súčasťou odovzdania. Funkčné testy a porovnanie celého payloadu sú potrebné, pretože samotné zníženie latencie by neodhalilo stratu historických odpovedí alebo nesprávny výber posledného pokusu.')
controls={ph:readj(P/'results'/ph/'controls.json')['metrics'] for ph in ['before','after']}
table(['Kontrolný scenár, 1 VU / 20 iterácií','p95 B / A (ms)'],[['Zoznam kurzov',f(controls['before']['course_list_ms']['values']['p(95)'],2)+' / '+f(controls['after']['course_list_ms']['values']['p(95)'],2)],['Odoslanie testu',f(controls['before']['submit_ms']['values']['p(95)'],2)+' / '+f(controls['after']['submit_ms']['values']['p(95)'],2)]],[300,155])
para('Kontroly používajú kópie databázy a reálny študentský login/CSRF; 40 odoslaní malo správny uložený výsledok. p95 odoslania mierne vzrástol. Tieto nezmenené endpointy sú kontrolou, nie dôkazom ich zrýchlenia.')
page();heading('7 AFTER a kvantitatívne porovnanie')
table(['Scenár','p95 BEFORE','p95 AFTER','Zníženie %','RPS B / A'],[[r['test'],f(row('before',r['test'])['p(95)']),f(r['p(95)']),f((row('before',r['test'])['p(95)']-r['p(95)'])/row('before',r['test'])['p(95)']*100),f(row('before',r['test'])['rps'],2)+' / '+f(r['rps'],2)] for r in rows if r['phase']=='after'],[75,95,95,85,105])
chart('p95.png');chart('throughput.png')
para('Zníženie latencie = (BEFORE − AFTER) / BEFORE × 100 %. Pri stress teste ostal p95 na hranici timeoutu: chyby klesli z 58,93 % na 21,36 %, počet prerušených iterácií sa zmenil z 29 na 59. AFTER load mal 0 prerušených iterácií. Celkové percentily stress zahŕňajú rôzne záťaže; podrobné hodnoty jednotlivých behov sú zachované v exportoch.')
page();heading('8 Dosiahnuté výsledky a zhodnotenie')
measured=[
 ['p95 pri 20 VUs (PR-01)']+[r.split(';')[0] for r in req[0][1:]],
 ['p95 pri 50 VUs (PR-02)']+[r.split(';')[0] for r in req[1][1:]],
 ['Load: chyby / prerušenia (PR-03)']+[f(row(ph,'load-3')['error']*100,2)+' % / '+str(row(ph,'load-3')['interrupted_iterations']) for ph in ['before','after']],
 ['Soak: posledná / prvá tretina (PR-04)']+[f(ratios[ph],3)+'×' for ph in ['before','after']],
 ['Počty: študenti / pokusy / posledné výsledky (PR-05)','1000 / 5000 / 1000','1000 / 5000 / 1000'],
]
table(['Ukazovateľ','BEFORE','AFTER'],measured,[175,140,140]);chart('stress.png')
para('Tabuľka sumarizuje namerané hodnoty ukazovateľov definovaných v kapitole 2. Časy pri 20 a 50 VUs vychádzajú zo stabilných 30 s okien. Graf zahŕňa aj 10 s na dokončenie požiadaviek, preto sa jeho hodnoty môžu líšiť. Pomer tretín vyjadruje vývoj odozvy počas trojminútového testu stability.')
para('Dosiahnuté prínosy: p95 celého load behu na LARGE sa znížil približne o 56 %. Po úprave tento beh zaznamenal 0 % chýb a 0 prerušených iterácií. Úplný výstup analytiky zostal zhodný vo všetkých troch kurzoch. Správnosť aplikácie podporuje 51 úspešných automatizovaných testov s 244 kontrolami.')
para('Podmienky merania: server a generátor záťaže bežali na jednom počítači, s jedným PHP workerom, databázou SQLite a spoločnou session učiteľa. Každá konfigurácia mala jeden platný HTTP beh pred úpravou a po nej. Výsledky opisujú túto lokálnu konfiguráciu; stabilita bola sledovaná počas 180 sekúnd.')
para('Záver: testovanie pomohlo identifikovať opakované spracovanie údajov, vybrať konkrétnu úpravu a zmerať jej účinok. Porovnanie pred úpravou a po nej preukázalo zrýchlenie pri zachovaní správnosti výsledkov. Projekt ukazuje praktický prínos výkonnostného testovania pri zlepšovaní aplikácie.')
page()
heading('Použité zdroje / Literatúra')
sources=[
'[1] Grafana Labs. k6: Metrics. https://grafana.com/docs/k6/latest/using-k6/metrics/',
'[2] Grafana Labs. k6: Scenarios. https://grafana.com/docs/k6/latest/using-k6/scenarios/',
'[3] Grafana Labs. k6: Cookies. https://grafana.com/docs/k6/latest/using-k6/cookies/',
'[4] PHP Group. Built-in web server. https://www.php.net/manual/en/features.commandline.webserver.php',
'[5] Laravel. Database: Listening for query events; Eloquent: Mutators and Casting. https://laravel.com/docs/12.x/database ; https://laravel.com/docs/12.x/eloquent-mutators',
]
for s in sources:story.append(Paragraph(html.escape(s),styles['CellSK']));story.append(Spacer(1,7))
pdf=D/'KS_Semestralny_projekt_Tema14_Mykhailo_Adamenko.pdf'
SimpleDocTemplate(str(pdf),pagesize=A4,rightMargin=70,leftMargin=70,topMargin=55,bottomMargin=45,title='Téma 14 – Výkonnostné testovanie',author='Mykhailo Adamenko').build(story,onFirstPage=footer,onLaterPages=footer)
import sys
if '--report-only' in sys.argv:
 print(pdf)
 sys.exit(0)
# Presentation: 12 slides, about 12–14 minutes including explanation of limitations.
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
slides=[
('Téma 14 – Výkonnostné testovanie',['E-learning na webovej platforme','Mykhailo Adamenko · FEI STU · 2026'],'Predstavte existujúcu aplikáciu a cieľ experimentu. 40 s.',None),
('Cieľ',['Zmerať učiteľskú analytiku pri raste dát','Nájsť príčinu, vykonať úpravu, overiť správnosť','Skutočné BEFORE / AFTER, bez odhadnutých výsledkov'],'Odlišujte zlepšenie od splnenia konkrétnej požiadavky. 50 s.',None),
('Architektúra a rozsah',['k6 → HTTP / Inertia → Laravel → SQLite','GET /teacher/courses/{id}/manage','Vue rendering nie je súčasťou merania'],'Vysvetlite roly a význam analytiky; scope je jeden endpoint. 60 s.',None),
('Požiadavky',['20 VUs: p95 < 800 ms; 50 VUs: p95 < 1500 ms','Chyby < 1 %; správne počty pri 1000 študentoch','Soak: posledná/prvá tretina p95 ≤ 1,25'],'Požiadavky boli definované vopred; ide o lokálne akceptačné ciele. 60 s.',None),
('Dataset a nástroje',['50 / 250 / 1000 študentov; 20 lekcií na kurz','5 pokusov na študenta, 10 otázok; seed 202627','k6, DB::listen, PHP časovače, PHPUnit'],'Samostatná SQLite databáza chráni pôvodné dáta. 60 s.',None),
('Load / Stress / Soak',['Load: 20 VUs, 5 + 30 + 5 s','Stress: 10 → 20 → 50 → 100 VUs','Soak: 5 VUs, 180 s; pripravená 1800 s verzia'],'Reálny login a CSRF; spoločná učiteľská session; jeden worker. 70 s.',None),
('BEFORE',['LARGE load p95: '+f(before['p(95)'])+' ms','Stress: timeouty, 29 prerušených iterácií','Fronta ovplyvnila nasledujúci setup'],'Vysvetlite zachované neúspešné behy a restart pred platným soak. 70 s.',None),
('Identifikovaná príčina',['CLI profil: '+f(prof['before']['ms'])+' ms; SQL '+f(prof['before']['sql_ms'])+' ms','Opakované spracovanie answers vo vnorenom cykle','Nepotrebné modely a opakované triedenie'],'Ukážte súvislosť profilu s vybranými opatreniami. 75 s.',None),
('Úprava a funkčná ekvivalencia',['Odpovede dekódované raz; jednoduchšie vnútorné cykly','Len potrebné stĺpce, bez zbytočného eager loading','Regresné testy a identický analytics JSON na 3 datasetoch'],'Žiadna cache, zmena schémy alebo odstránenie histórie. 75 s.',None),
('BEFORE / AFTER',['LARGE load: zníženie p95 o '+f(gain)+' %'],'Vyložte graf aj prípadné nesplnené limity. 80 s.','p95.png'),
('Splnenie požiadaviek',[r[0]+': BEFORE '+r[1]+' | AFTER '+r[2] for r in req],'PR-04 je len krátka verzia; PR-05 je správnosť. 80 s.',None),
('Záver a obmedzenia',['Profilovanie viedlo ku konkrétnej overenej úprave','Lokálny jeden worker ≠ produkčná kapacita','Ďalej: opakované behy, produkčný server, plný soak','Zdroje a reprodukcia: PDF + performance-testing/README.md'],'Uzavrite prínos a ukážte umiestnenie artefaktov. Zdroje: Grafana k6 Metrics/Scenarios/Cookies; Laravel Database/Casting; PHP Built-in web server; zadanie a bakalárska práca. 60 s.',None),
]
md=[]
for i,(title,bullets,notes,picture) in enumerate(slides,1):
 slide=prs.slides.add_slide(prs.slide_layouts[6]);box=slide.shapes.add_textbox(Inches(.7),Inches(.45),Inches(12),Inches(.8));tf=box.text_frame;tf.text=title;tf.paragraphs[0].font.size=Pt(30);tf.paragraphs[0].font.bold=True;tf.paragraphs[0].font.color.rgb=RGBColor.from_string('176C83')
 box=slide.shapes.add_textbox(Inches(.8),Inches(1.55),Inches(11.8),Inches(4.8));tf=box.text_frame;tf.word_wrap=True
 for j,b in enumerate(bullets):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=b;p.font.size=Pt(22 if len(bullets)<5 else 19);p.space_after=Pt(22)
 if picture:slide.shapes.add_picture(str(P/'reports'/picture),Inches(1.6),Inches(2.4),width=Inches(10))
 foot=slide.shapes.add_textbox(Inches(.8),Inches(7),Inches(11.8),Inches(.3));foot.text_frame.text=f'Mykhailo Adamenko · KS 2026/2027                                         {i}/12';foot.text_frame.paragraphs[0].font.size=Pt(11)
 slide.notes_slide.notes_text_frame.text=notes
 md.append('## '+str(i)+'. '+title+'\n\n'+'\n'.join('- '+b for b in bullets)+'\n\nPoznámky: '+notes)
prs.save(str(D/'KS_Tema14_Prezentacia_Mykhailo_Adamenko.pptx'))
(D/'presentation.md').write_text('\n\n'.join(md),encoding='utf-8')
print(pdf)
# Keep the redesigned, editable presentation as the canonical version.
import runpy
runpy.run_path(str(P/'scripts/build_presentation.py'),run_name='__main__')
